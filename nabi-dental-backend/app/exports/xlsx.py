from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.schemas.report import ClinicReport, ConstructionReport, HomeReport
from app.utils.money import format_money, format_quantity

HEADER_FILL = PatternFill("solid", fgColor="7C3AED")
SOFT_FILL = PatternFill("solid", fgColor="F3E8FF")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri")
TITLE_FONT = Font(bold=True, size=16, color="1F2937", name="Calibri")
SECTION_FONT = Font(bold=True, size=12, color="7C3AED", name="Calibri")
LABEL_FONT = Font(bold=True, name="Calibri", color="1F2937")
THIN = Border(
    left=Side(style="thin", color="E5E7EB"),
    right=Side(style="thin", color="E5E7EB"),
    top=Side(style="thin", color="E5E7EB"),
    bottom=Side(style="thin", color="E5E7EB"),
)


def render_clinic_xlsx(report: ClinicReport, *, clinic_name: str) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _write_header(summary, clinic_name, "Clinic Expense Report", report.period.from_date, report.period.to_date, report.currency)
    _write_kv(
        summary,
        7,
        [
            ("Total income", format_money(report.total_income)),
            ("Total clinic expenses", format_money(report.total_expenses)),
            ("Profit / loss", format_money(report.profit)),
            ("Income entries", report.income_count),
            ("Expense entries", report.expense_count),
        ],
    )
    _write_named_sheet(workbook, "Income by treatment", ["Treatment", "Amount"], [(item.name, item.amount) for item in report.income_by_treatment])
    _write_named_sheet(workbook, "Expenses by category", ["Category", "Amount"], [(item.name, item.amount) for item in report.expenses_by_category])
    _write_table_sheet(
        workbook,
        "Income transactions",
        ["Date", "Treatment", "Notes", "Amount"],
        [[str(item.entry_date), item.name, item.detail or "", item.amount] for item in report.income_lines],
        money_cols={4},
    )
    _write_table_sheet(
        workbook,
        "Clinic expenses",
        ["Date", "Category", "Notes", "Amount"],
        [[str(item.entry_date), item.name, item.detail or "", item.amount] for item in report.expense_lines],
        money_cols={4},
    )
    _write_timeline_sheet(workbook, report.timeline, include_income=True)
    return _save(workbook)


def render_home_xlsx(report: HomeReport, *, clinic_name: str) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _write_header(summary, clinic_name, "Home Expense Report", report.period.from_date, report.period.to_date, report.currency)
    used = f"{report.percentage_used:.2f}" if report.percentage_used is not None else "n/a"
    _write_kv(
        summary,
        7,
        [
            ("Total spending", format_money(report.total_expenses)),
            ("Budget", format_money(report.budget)),
            ("Remaining", format_money(report.remaining)),
            ("Percentage used", used),
            ("Transactions", report.expense_count),
        ],
    )
    _write_named_sheet(workbook, "Spending by category", ["Category", "Amount"], [(item.name, item.amount) for item in report.expenses_by_category])
    _write_table_sheet(
        workbook,
        "Home expenses",
        ["Date", "Category", "Notes", "Amount"],
        [[str(item.entry_date), item.name, item.detail or "", item.amount] for item in report.expense_lines],
        money_cols={4},
    )
    _write_timeline_sheet(workbook, report.timeline, include_income=False)
    return _save(workbook)


def render_construction_xlsx(report: ConstructionReport, *, clinic_name: str) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _write_header(
        summary,
        clinic_name,
        "Construction Expense Report",
        report.period.from_date,
        report.period.to_date,
        report.currency,
    )
    _write_kv(
        summary,
        7,
        [
            ("Total construction spend", format_money(report.total_expenses)),
            ("Purchases", report.purchase_count),
            ("Categories", len(report.expenses_by_category)),
            ("Suppliers", len(report.expenses_by_supplier)),
        ],
    )
    _write_named_sheet(
        workbook,
        "By category",
        ["Category", "Amount"],
        [(item.name, item.amount) for item in report.expenses_by_category],
    )
    _write_table_sheet(
        workbook,
        "By material",
        ["Material", "Quantity", "Unit", "Total"],
        [[item.name, format_quantity(item.quantity), item.unit, item.amount] for item in report.spending_by_material],
        money_cols={4},
        number_cols={2},
    )
    _write_named_sheet(
        workbook,
        "By supplier",
        ["Supplier", "Amount"],
        [(item.name, item.amount) for item in report.expenses_by_supplier],
    )
    _write_table_sheet(
        workbook,
        "Purchases",
        ["Date", "Material", "Category", "Quantity", "Unit", "Unit price", "Total", "Supplier"],
        [
            [
                str(item.purchase_date),
                item.material_name,
                item.category_name,
                format_quantity(item.quantity),
                item.unit,
                item.unit_price,
                item.amount,
                item.supplier or "",
            ]
            for item in report.purchases
        ],
        money_cols={6, 7},
        number_cols={4},
    )
    _write_timeline_sheet(workbook, report.timeline, include_income=False)
    return _save(workbook)


def _write_header(sheet, clinic_name, title, from_date, to_date, currency) -> None:
    sheet["A1"] = clinic_name
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = title
    sheet["A2"].font = SECTION_FONT
    sheet["A3"] = "From"
    sheet["B3"] = str(from_date)
    sheet["A4"] = "To"
    sheet["B4"] = str(to_date)
    sheet["A5"] = "Currency"
    sheet["B5"] = currency
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 22


def _write_kv(sheet, start_row: int, rows: list[tuple]) -> None:
    for offset, (label, value) in enumerate(rows):
        row = start_row + offset
        sheet.cell(row, 1, label).font = LABEL_FONT
        sheet.cell(row, 1).fill = SOFT_FILL
        sheet.cell(row, 2, value)


def _write_named_sheet(workbook, title: str, headers: list[str], rows: list[tuple]) -> None:
    _write_table_sheet(workbook, title, headers, [list(row) for row in rows], money_cols={2})


def _write_timeline_sheet(workbook, points, *, include_income: bool) -> None:
    if include_income:
        _write_table_sheet(
            workbook,
            "Timeline",
            ["Period", "Income", "Expenses", "Profit"],
            [
                [item.period, item.income, item.expenses, item.profit]
                for item in points
            ],
            money_cols={2, 3, 4},
        )
        return
    _write_table_sheet(
        workbook,
        "Timeline",
        ["Period", "Amount"],
        [[item.period, item.expenses] for item in points],
        money_cols={2},
    )


def _write_table_sheet(
    workbook,
    title: str,
    headers: list[str],
    rows: list[list],
    money_cols: set[int] | None = None,
    number_cols: set[int] | None = None,
) -> None:
    sheet = workbook.create_sheet(title[:31])
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(1, index, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN
        sheet.column_dimensions[get_column_letter(index)].width = 18
    money_cols = money_cols or set()
    number_cols = number_cols or set()
    for row_index, row in enumerate(rows, start=2):
        for col_index, value in enumerate(row, start=1):
            cell = sheet.cell(row_index, col_index, _excel_value(value, col_index in money_cols))
            cell.border = THIN
            if col_index in money_cols and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"
            elif col_index in number_cols:
                cell.number_format = "0.###"
    if not rows:
        sheet.cell(2, 1, "No records in this period.")


def _excel_value(value, as_money: bool):
    if value is None:
        return None
    if as_money:
        try:
            return float(format_money(value))
        except Exception:
            return str(value)
    return value


def _save(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
